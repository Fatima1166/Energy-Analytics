"""scripts/inspect_deck.py -- dump the structure/design of a .pptx for editing."""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

def inches(value) -> float:
    return round(Emu(value).inches, 2) if value is not None else -1.0


def colour_of(run) -> str:
    try:
        rgb = run.font.color.rgb
        return f"#{rgb}" if rgb is not None else "-"
    except Exception:
        return "theme"


def main(path: str) -> int:
    prs = Presentation(path)
    print(f"file            : {path}")
    print(f"slide size      : {inches(prs.slide_width)} x {inches(prs.slide_height)} in "
          f"({len(prs.slides)} slides)")
    print(f"layouts         : {[l.name for l in prs.slide_layouts]}")
    print()
    for index, slide in enumerate(prs.slides, start=1):
        print("=" * 92)
        print(f"SLIDE {index}   layout={slide.slide_layout.name}")
        for shape in slide.shapes:
            kind = str(shape.shape_type).split(" ")[0]
            pos = f"({inches(shape.left)},{inches(shape.top)}) {inches(shape.width)}x{inches(shape.height)}"
            print(f"  - {shape.shape_type} '{shape.name}' {pos}")
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs)
                    if not text.strip():
                        continue
                    meta = " ".join(
                        f"[{r.font.size.pt if r.font.size else '-'}pt"
                        f"{' b' if r.font.bold else ''} {colour_of(r)}]"
                        for r in para.runs[:2])
                    print(f"      {meta} {text[:130]}")
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.replace("\n", " ")[:28] for c in row.cells]
                    print("      | " + " | ".join(cells) + " |")
            if shape.shape_type == 13:      # picture
                print(f"      (picture {shape.image.filename or shape.image.content_type})")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"  notes: {slide.notes_slide.notes_text_frame.text[:120]}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1
                          else r"C:\Users\PC\Downloads\FUI_Energy_Analytics.pptx"))
