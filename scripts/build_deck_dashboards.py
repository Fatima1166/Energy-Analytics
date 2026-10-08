"""
scripts/build_deck_dashboards.py -- append dashboard "how it works" slides.

Adds six slides to the existing FUI Energy Analytics deck, in the deck's own
design language (navy #0B1F3A titles, white cards, navy table headers):

    14  The Three Dashboards          -- what each one answers
    15  How Every Dashboard Works     -- the same four steps, as a flow
    16  Dashboard 1: Renewable Scenarios (Q2)
    17  Dashboard 2: Target Bottlenecks (Q3)
    18  Dashboard 3: Forecasted Demand (Q4)
    19  Reported vs Modelled          -- the caveats before quoting numbers

The original file is backed up next to it before the in-place save.
"""

from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

DECK = Path(r"C:\Users\PC\Downloads\FUI_Energy_Analytics.pptx")

# ── the deck's own palette ──────────────────────────────────────────────────
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GREY = RGBColor(0x4A, 0x55, 0x68)
GREEN = RGBColor(0x0E, 0x9F, 0x8E)
RED = RGBColor(0xD6, 0x39, 0x3F)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF4, 0xF7, 0xFB)
BORDER = RGBColor(0xD6, 0xDE, 0xE9)

TITLE_SIZE = 34
BODY_SIZE = 14


# ═════════════════════════════ building blocks ══════════════════════════════

def add_slide(prs: Presentation):
    """A blank slide on the deck's DEFAULT layout, with no inherited placeholders."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            shape._element.getparent().remove(shape._element)
    return slide


def text(slide, x, y, w, h, value, size=BODY_SIZE, bold=False, colour=GREY,
         align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    return box


def slide_title(slide, value: str, colour=NAVY):
    return text(slide, 0.5, 0.35, 9.0, 0.7, value, size=TITLE_SIZE, bold=True,
                colour=colour)


def subtitle(slide, value: str):
    return text(slide, 0.5, 1.05, 9.0, 0.35, value, size=BODY_SIZE, colour=GREY)


def footnote(slide, value: str, colour=NAVY, y=4.98):
    return text(slide, 0.5, y, 9.0, 0.5, value, size=BODY_SIZE, colour=colour)


def card(slide, x, y, w, h, heading=None, lines=(), number=None,
         heading_size=16, body_size=13, accent=NAVY, heading_colour=NAVY,
         fill=CARD):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                   Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.text_frame.text = ""

    text_x = x + 0.22
    text_w = w - 0.44
    if number is not None:
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2),
                                       Inches(y + 0.18), Inches(0.42), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = accent
        badge.line.fill.background()
        badge.shadow.inherit = False
        frame = badge.text_frame
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = str(number)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = WHITE
        text_x = x + 0.72
        text_w = w - 0.94

    cursor = y + 0.18
    if heading:
        text(slide, text_x, cursor, text_w, 0.35, heading, size=heading_size,
             bold=True, colour=heading_colour)
        cursor += 0.38
    for line in lines:
        box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(cursor),
                                       Inches(w - 0.44), Inches(0.6))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        para = frame.paragraphs[0]
        run = para.add_run()
        run.text = line
        run.font.size = Pt(body_size)
        run.font.color.rgb = GREY
        cursor += 0.26 if len(line) < 60 else 0.42
    return shape


def stat(slide, x, y, w, h, value, label, accent=NAVY, note=None, value_size=24):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                   Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    text(slide, x + 0.22, y + 0.14, w - 0.44, 0.45, value, size=value_size, bold=True,
         colour=accent)
    text(slide, x + 0.22, y + 0.58, w - 0.44, 0.3, label, size=12, colour=GREY)
    if note:
        text(slide, x + 0.22, y + 0.84, w - 0.44, 0.3, note, size=11, colour=GREY)
    return shape


def table(slide, data, x, y, w, h, col_widths=None, font=BODY_SIZE,
          label_col=True):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w),
                                   Inches(h))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    if col_widths:
        total = sum(col_widths)
        for index, width in enumerate(col_widths):
            tbl.columns[index].width = Emu(int(Inches(w) * width / total))
    for r in range(rows):
        tbl.rows[r].height = Inches(h / rows)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else WHITE
            frame = cell.text_frame
            frame.word_wrap = True
            para = frame.paragraphs[0]
            run = para.add_run()
            run.text = str(data[r][c])
            run.font.size = Pt(font)
            if r == 0:
                run.font.bold = True
                run.font.color.rgb = WHITE
            elif c == 0 and label_col:
                run.font.bold = True
                run.font.color.rgb = NAVY
            else:
                run.font.color.rgb = GREY
    return shape


def arrow(slide, x, y, w=0.3, h=0.26):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


_NOTES_SP_XML = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="9" name="Notes Placeholder"/>'
    '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/>'
    '<a:ext cx="5486400" cy="3600450"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
)


def notes(slide, value: str):
    """
    Speaker notes.  This deck's layout carries no notes placeholder, so one is
    injected into the notes slide before writing; if anything about that fails
    the slide is still usable and the build continues.
    """
    try:
        notes_slide = slide.notes_slide
        if notes_slide.notes_text_frame is None:
            from lxml import etree
            notes_slide._element.spTree.append(etree.fromstring(_NOTES_SP_XML))
        notes_slide.notes_text_frame.text = value
    except Exception as exc:                     # pragma: no cover - defensive
        print(f"  (notes skipped on this slide: {exc})")


# ═══════════════════════════════ the slides ═════════════════════════════════

def slide_three_dashboards(prs):
    slide = add_slide(prs)
    slide_title(slide, "The Three Dashboards")
    subtitle(slide, "One database, one unit system, three different questions.")
    table(slide, [
        ["Dashboard", "Question it answers", "What you get"],
        ["Q2  Renewable\nScenarios",
         "What are the potential outcomes of installing renewable energy sources?",
         "RE share, capacity, CO\u2082 and LNG savings to 2035 under 5 scenarios"],
        ["Q3  Target\nBottlenecks",
         "What are the bottlenecks in achieving planned targets?",
         "4 KPIs scored against benchmarks and ranked worst-first"],
        ["Q4  Forecasted\nDemand",
         "What is forecasted energy demand?",
         "Demand by fuel, sector and province to 2035, with drivers and a backtest"],
    ], 0.5, 1.5, 9.0, 2.9, col_widths=[1.5, 3.4, 4.1], font=12)
    footnote(slide, "Pick a dashboard in the sidebar of the Streamlit app "
                    "(http://localhost:8501) \u2014 the page, the slicers and the "
                    "units all change with it.")
    notes(slide, "Three dashboards, one database. Say the question first, then the "
                 "output \u2014 that is the whole structure of the app.")


def slide_pipeline(prs):
    slide = add_slide(prs)
    slide_title(slide, "How Every Dashboard Works")
    subtitle(slide, "The same four steps behind all three dashboards.")
    steps = [
        ("1", "MySQL", "35 tables in a star schema, every value in one base unit."),
        ("2", "SQL", "core/queries.py pulls the raw rows, cached for 10 minutes."),
        ("3", "Engine", "Models the future: trends, elasticities, scenarios, backtest."),
        ("4", "Streamlit", "KPI cards and Plotly charts in the unit you pick."),
    ]
    x = 0.5
    for index, (number, heading, body) in enumerate(steps):
        card(slide, x, 1.55, 1.95, 2.6, heading=heading, lines=[body],
             number=number, heading_size=13, body_size=11, accent=NAVY)
        if index < len(steps) - 1:
            arrow(slide, x + 1.99, 2.7)
        x += 2.35
    footnote(slide, "Data stays in base units (TOE, GWh, MW, BCF, t) the whole way; "
                    "the sidebar unit box converts only what is drawn on screen.",
             y=4.3)
    footnote(slide, "Nothing is hard-coded in the charts: move a slicer and the "
                    "engine recomputes the numbers.", colour=GREEN, y=4.7)
    notes(slide, "If someone asks 'how does it work', these four steps are the answer: "
                 "database, query, engine, page.")


def slide_q2(prs):
    slide = add_slide(prs)
    slide_title(slide, "Dashboard 1 \u2014 Renewable Scenarios")
    subtitle(slide, "Question 2 \u2014 what happens if we install renewables at "
                    "different speeds?")
    card(slide, 0.5, 1.55, 4.5, 3.2, heading="How it works", lines=[
        "\u2022  Solar and wind grow at the scenario rate (2% to 25% a year) from the last reported year",
        "\u2022  Fossil generation closes the balance: demand \u2212 renewables \u2212 nuclear",
        "\u2022  CO\u2082 = coal and gas generation \u00d7 emission factor (0.90 / 0.40 tCO\u2082/MWh)",
        "\u2022  LNG saving = gas generation avoided \u00d7 $75m per TWh",
    ], heading_size=17, body_size=12)
    stat(slide, 5.2, 1.55, 4.3, 1.0, "73.5%", "renewable share of generation in 2035",
         accent=GREEN, note="Accelerated scenario (15% a year)")
    stat(slide, 5.2, 2.65, 4.3, 1.0, "19.7 Mt", "CO\u2082 left in 2035",
         accent=NAVY, note="down from 39.5 Mt reported in 2024")
    stat(slide, 5.2, 3.75, 4.3, 1.0, "$21.6 bn", "cumulative LNG savings, 2025\u20132035",
         accent=ORANGE, note="gas imports avoided by clean generation")
    footnote(slide, "Five scenarios are compared at the same target year, and the "
                    "growth-rate slider can override the active one.")
    notes(slide, "Key point: the scenario only changes how fast renewables are built. "
                 "Everything else follows from the balance.")


def slide_q3(prs):
    slide = add_slide(prs)
    slide_title(slide, "Dashboard 2 \u2014 Target Bottlenecks")
    subtitle(slide, "Question 3 \u2014 are we hitting the planned targets, and where "
                    "does it hurt most?")
    table(slide, [
        ["KPI", "2035", "Benchmark", "Gap", "Verdict"],
        ["Energy losses %", "27.16", "28.23", "+1.07", "on track"],
        ["Import dependence %", "38.26", "0.00", "\u221238.26", "severe"],
        ["Renewable share %", "50.05", "30.00", "+20.05", "exceeding"],
        ["Gas self-sufficiency %", "47.80", "122.37", "\u221274.57", "#1 bottleneck"],
    ], 0.5, 1.55, 9.0, 2.5, col_widths=[3.0, 1.2, 1.5, 1.2, 2.1], font=12)
    footnote(slide, "Benchmarks live in dim_target and are direction-aware, so a "
                    "positive gap always means the target is met.")
    footnote(slide, "Ranking weights the size of the miss and the shortfall relative "
                    "to the benchmark, so a small target missed badly is not hidden.",
             colour=GREEN, y=4.45)
    notes(slide, "The verdict column is what the ranking uses: worst miss first, so a "
                 "reader sees the bottleneck before the detail.")


def slide_q4(prs):
    slide = add_slide(prs)
    slide_title(slide, "Dashboard 3 \u2014 Forecasted Demand")
    subtitle(slide, "Question 4 \u2014 what is forecasted energy demand, and what "
                    "drives it?")
    card(slide, 0.5, 1.55, 4.5, 3.2, heading="Six objectives, eight elements", lines=[
        "1  Reported demand, 1999\u20132024",
        "2  Low / Base / High scenarios to 2035",
        "3  Drivers: GDP, GDP per head, population",
        "4  Sector and province breakdown",
        "5  Supply gap and import dependency",
        "6  Backtest: MAPE, RMSE, R\u00b2",
    ], heading_size=17, body_size=12)
    stat(slide, 5.2, 1.55, 4.3, 1.0, "76.6 \u2192 95.2 MTOE",
         "total final energy, 2024 \u2192 2035 Base", accent=NAVY)
    stat(slide, 5.2, 2.65, 4.3, 1.0, "175,440 \u2192 300,586 GWh",
         "grid electricity demand, 2024 \u2192 2035 Base", accent=GREEN, value_size=21)
    stat(slide, 5.2, 3.75, 4.3, 1.0, "1.04", "income elasticity of demand",
         accent=ORANGE, note="R\u00b2 0.92 \u00b7 population +1.6% a year")
    footnote(slide, "Growth = population + elasticity \u00d7 GDP per head + residual; "
                    "Low/Base/High = GDP \u00b1 1.5 pp. Hold-out MAPE 3.7\u201323.3%.")
    notes(slide, "Point at the elasticity: it is the bridge between the economy and "
                 "demand, and it is what makes the scenarios differ.")


def slide_caveats(prs):
    slide = add_slide(prs)
    slide_title(slide, "Reported vs Modelled")
    subtitle(slide, "Read this before quoting any number after the last reported year.")
    table(slide, [
        ["Dashboard", "Reported to", "Modelled", "Main caveat"],
        ["Q2 Renewables", "2024", "2025\u20132035",
         "demand growth is an assumption; capacity table ends 2018"],
        ["Q3 Bottlenecks", "losses 2019, imports 2018", "to 2035",
         "trends extended; gas growth floored at \u00b12% a year"],
        ["Q4 Demand", "final energy 2024, gas 2018", "2025\u20132035",
         "2019 coverage change handled with a dummy; residual fades by 2035"],
    ], 0.5, 1.5, 9.0, 2.9, col_widths=[1.7, 2.2, 1.3, 3.8], font=11)
    footnote(slide, "Reported years are drawn solid and modelled years dashed on every "
                    "chart, so a projection is never mistaken for a measurement.",
             colour=NAVY)
    notes(slide, "Say the caveats out loud \u2014 it is what separates a forecast from a "
                 "guess, and the charts label it too.")


OUR_TITLES = {
    "The Three Dashboards",
    "How Every Dashboard Works",
    "Dashboard 1 \u2014 Renewable Scenarios",
    "Dashboard 2 \u2014 Target Bottlenecks",
    "Dashboard 3 \u2014 Forecasted Demand",
    "Reported vs Modelled",
}


def drop_our_slides(prs: Presentation) -> int:
    """Remove any previously-added dashboard slides so the build is idempotent."""
    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    removed = 0
    # walk by index so the relationship id can be dropped as well
    for index in range(len(entries) - 1, -1, -1):
        slide = prs.slides[index]
        first_text = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                first_text = shape.text_frame.text.strip()
                break
        if first_text in OUR_TITLES:
            rel_id = entries[index].get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            prs.part.drop_rel(rel_id)
            id_list.remove(entries[index])
            removed += 1
    return removed


def main() -> int:
    if not DECK.exists():
        print(f"deck not found: {DECK}")
        return 1

    existing = sorted(DECK.parent.glob(f"{DECK.stem}.backup-*{DECK.suffix}"))
    if existing:
        print(f"backup          : {existing[0].name} (kept)")
    else:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = DECK.with_name(f"{DECK.stem}.backup-{stamp}{DECK.suffix}")
        shutil.copy2(DECK, backup)
        print(f"backup          : {backup.name}")

    prs = Presentation(str(DECK))
    print(f"slides before   : {len(prs.slides)}")
    dropped = drop_our_slides(prs)
    if dropped:
        print(f"replaced        : {dropped} previously added dashboard slides")

    slide_three_dashboards(prs)
    slide_pipeline(prs)
    slide_q2(prs)
    slide_q3(prs)
    slide_q4(prs)
    slide_caveats(prs)

    prs.save(str(DECK))
    print(f"slides after    : {len(Presentation(str(DECK)).slides)}")
    print(f"saved           : {DECK}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
