"""
scripts/make_render_copy.py -- copy the deck at a 13.33x7.5in page size.

The bundled LibreOffice renderer mis-scales this deck's non-default 10x5.62in
page (it renders every slide ~1.33x too large and crops the right edge, on the
original slides too).  Enlarging only the PAGE, in a throwaway copy, keeps every
shape's EMU position intact while letting the renderer produce a faithful image,
so text fit inside cards and table cells can be checked visually.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def main(src: str, dst: str) -> int:
    shutil.copy2(src, dst)
    prs = Presentation(dst)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.save(dst)
    print(f"render copy: {dst} "
          f"({len(Presentation(dst).slides)} slides at 13.333 x 7.5 in)")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main(
        sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\PC\Downloads\FUI_Energy_Analytics.pptx",
        sys.argv[2] if len(sys.argv) > 2 else r"D:\Fatima Projects\energy_dashboard\deck_render_copy.pptx"))
