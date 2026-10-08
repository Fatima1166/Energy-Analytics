"""scripts/probe_deck_table.py -- report how the existing deck styles tables."""

from __future__ import annotations

import sys

from pptx import Presentation


def describe_cell(cell, label: str):
    tf = cell.text_frame
    para = tf.paragraphs[0]
    run = para.runs[0] if para.runs else None
    fill = "-"
    try:
        if cell.fill.type is not None:
            fill = f"{cell.fill.type}"
            if cell.fill.type == 1:      # solid
                fill += f" {cell.fill.fore_color.rgb}"
    except Exception as exc:             # pragma: no cover
        fill = f"err {exc}"
    print(f"    {label}: text={cell.text[:26]!r} size="
          f"{run.font.size.pt if run and run.font.size else '-'}pt "
          f"bold={run.font.bold if run else '-'} "
          f"color={run.font.color.rgb if run and run.font.color and run.font.color.type is not None else '-'} "
          f"fill={fill} align={para.alignment}")


def main(path: str) -> int:
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            print(f"SLIDE {idx}: table {shape.name}  rows={len(table.rows)} "
                  f"cols={len(table.columns)} first_row={table.first_row} "
                  f"banding={table.horz_banding}")
            xml = table._tbl.xml
            marker = "tblStyle"
            if marker in xml:
                start = xml.index(marker)
                print(f"    style: ...{xml[start:start + 60]}...")
            else:
                print("    style: (none - default)")
            describe_cell(table.cell(0, 0), "header[0,0]")
            if len(table.rows) > 1:
                describe_cell(table.cell(1, 0), "body[1,0]")
            print()
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1
                          else r"C:\Users\PC\Downloads\FUI_Energy_Analytics.pptx"))
