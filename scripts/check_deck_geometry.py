"""scripts/check_deck_geometry.py -- bounding boxes per slide (renderer-independent)."""

from __future__ import annotations

import sys

from pptx import Presentation
from pptx.util import Emu


def inches(value) -> float:
    return round(Emu(value).inches, 2) if value is not None else 0.0


def main(path: str, pages) -> int:
    prs = Presentation(path)
    print(f"slide size: {inches(prs.slide_width)} x {inches(prs.slide_height)} in "
          f"({len(prs.slides)} slides)")
    for index, slide in enumerate(prs.slides, start=1):
        if pages and index not in pages:
            continue
        lefts, rights, tops, bottoms = [], [], [], []
        widest = None
        for shape in slide.shapes:
            if shape.left is None:
                continue
            left, right = inches(shape.left), inches(shape.left) + inches(shape.width)
            top, bottom = inches(shape.top), inches(shape.top) + inches(shape.height)
            lefts.append(left); rights.append(right)
            tops.append(top); bottoms.append(bottom)
            if widest is None or right > widest[0]:
                widest = (right, shape.name, f"{left}->{right}")
        print(f"  slide {index:2d}: shapes={len(slide.shapes):2d} "
              f"x[{min(lefts, default=0):.2f}..{max(rights, default=0):.2f}] "
              f"y[{min(tops, default=0):.2f}..{max(bottoms, default=0):.2f}] "
              f"rightmost={widest}")
        over = [s.name for s in slide.shapes
                if s.left is not None and inches(s.left) + inches(s.width) > 9.75]
        if over:
            print(f"          OVERFLOW beyond 9.75in: {over}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = sys.argv[1:]
    path = args[0] if args else r"C:\Users\PC\Downloads\FUI_Energy_Analytics.pptx"
    pages = {int(p) for p in args[1:]} if len(args) > 1 else set()
    raise SystemExit(main(path, pages))
